"""Cambio 9 on real PostgreSQL: subcontracts, estimations, security and overview impact.

Run with D89_RUN_FINANCIAL_INTEGRATION=1 (disposable Docker services, never .env).
"""

import io
import time
from datetime import date
from decimal import Decimal
from uuid import uuid4

import jwt
import openpyxl
import psycopg
import pytest
from fastapi.testclient import TestClient
from test_financial_security_integration import (
    BOOTSTRAP,
    MIGRATIONS,
    isolated_services,  # noqa: F401 -- shared disposable pytest fixture
    seed,
)

from app.core.config import Settings, get_settings
from app.main import app

MIGRATION = MIGRATIONS / "202610010001_subcontract_module.sql"
BRIDGE = MIGRATIONS / "202610010002_estimation_expense_bridge.sql"
D = Decimal


@pytest.mark.parametrize("installation", ["fresh", "upgrade"])
def test_subcontract_lifecycle_security_and_overview(isolated_services, installation):  # noqa: F811
    services = isolated_services
    database = f"subcontract_{installation}"
    with psycopg.connect(services["dsn"], autocommit=True) as admin:
        admin.execute(
            psycopg.sql.SQL("create database {}").format(psycopg.sql.Identifier(database))
        )
    dsn = services["dsn"].removesuffix("postgres") + database
    legacy_id = uuid4()
    with psycopg.connect(dsn, autocommit=True) as connection:
        connection.execute(BOOTSTRAP)
        for path in sorted(MIGRATIONS.glob("*.sql")):
            if path == MIGRATION and installation == "upgrade":
                ids = seed(connection)
                # A contract and payment of the initial schema (free-text subcontractor).
                connection.execute(
                    """insert into public.subcontrato
                       (id, obra_id, subcontratista, concepto, alcance, monto_contratado)
                       values (%s, %s, 'Juan Pérez', 'Muros', 'Muros de block', 5000.5)""",
                    (legacy_id, ids["work"]),
                )
                connection.execute(
                    """insert into public.subcontrato_pago
                       (subcontrato_id, fecha, monto, creado_por)
                       values (%s, '2026-09-10', 1000, %s)""",
                    (legacy_id, ids["admin"]),
                )
            connection.execute(path.read_text())
        if installation == "fresh":
            ids = seed(connection)
        else:
            legacy = connection.execute(
                """select importe_contratado, folio, estado::text, subcontratista
                   from public.subcontrato where id = %s""",
                (legacy_id,),
            ).fetchone()
            assert legacy == (D("5000.5"), "SC-0001", "activo", "Juan Pérez")
        # Idempotent: reapplying keeps rows and never rewinds the folio sequence.
        before = connection.execute("select * from public.subcontrato order by id").fetchall()
        connection.execute(MIGRATION.read_text())
        assert connection.execute(
            "select * from public.subcontrato order by id"
        ).fetchall() == before

        # PostgREST roles can read at most; every write goes through FastAPI.
        for role in ("anon", "authenticated"):
            for table in ("subcontrato", "estimacion_subcontrato", "subcontrato_pago"):
                for privilege in ("INSERT", "UPDATE", "DELETE"):
                    assert not connection.execute(
                        "select has_table_privilege(%s, %s, %s)", (role, table, privilege)
                    ).fetchone()[0], (role, table, privilege)
        restrictive = connection.execute(
            """select count(*) from pg_policies where schemaname = 'public'
               and tablename in ('subcontrato', 'estimacion_subcontrato', 'subcontrato_pago')
               and permissive = 'RESTRICTIVE'"""
        ).fetchone()[0]
        assert restrictive == 9

        def headers(role):
            if role is None:
                return {}
            return {"Authorization": "Bearer " + jwt.encode(
                {"sub": str(ids[role]), "role": "authenticated", "aud": "authenticated",
                 "app_metadata": {"role": "admin" if role == "admin" else "operativo"},
                 "exp": int(time.time()) + 600},
                services["secret"], algorithm="HS256",
            )}

        settings = Settings(_env_file=None, database_url=dsn, jwt_secret=services["secret"])
        previous = app.dependency_overrides.copy()
        app.dependency_overrides[get_settings] = lambda: settings
        base = f"/api/v1/works/{ids['work']}/subcontracts"
        body = {
            "supplier_id": str(ids["supplier"]), "expense_item_id": str(ids["item"]),
            "expense_subitem_id": str(ids["subitem"]),
            "description": "Colocación de block en muros perimetrales",
            "contracted_amount": "10000", "retention_percent": "5",
        }
        try:
            with TestClient(app) as api:
                for role, expected in ((None, 401), ("operativo", 403), ("outsider", 403)):
                    assert api.post(base, json=body, headers=headers(role)).status_code == expected
                assert api.post(
                    base, json={**body, "supplier_id": str(uuid4())}, headers=headers("admin")
                ).status_code == 422
                created = api.post(base, json=body, headers=headers("admin"))
                assert created.status_code == 201, created.text
                contract = created.json()
                sid = contract["id"]
                assert contract["folio"] == ("SC-0002" if installation == "upgrade" else "SC-0001")
                assert contract["category"] == "MANO DE OBRA" and contract["state"] == "activo"
                assert contract["remaining_to_estimate"] == "10000.0000"
                listed = api.get(base, headers=headers("operativo"))
                assert listed.status_code == 200 and sid in [row["id"] for row in listed.json()]
                assert api.get(base, headers=headers("outsider")).status_code == 403

                est = f"/api/v1/subcontracts/{sid}/estimations"

                def post(payload, role="admin"):
                    return api.post(est, json={"estimated_on": "2026-09-20", **payload},
                                    headers=headers(role))

                def pay(estimation_id):
                    response = api.patch(f"{est}/{estimation_id}/status",
                                         json={"state": "pagado"}, headers=headers("admin"))
                    assert response.status_code == 200, response.text
                    return response.json()

                assert post({"kind": "anticipo", "gross_amount": "100"}, "operativo").status_code \
                    == 403
                advance = post({"kind": "anticipo", "gross_amount": "1000"})
                assert advance.status_code == 201, advance.text
                assert advance.json()["folio"] == "EST-01"
                assert advance.json()["net_amount"] == "1000.0000"
                # An unpaid advance cannot be amortized yet.
                assert post({"kind": "avance", "gross_amount": "4000",
                             "advance_amortization": "500"}).status_code == 422
                expenses_before = connection.execute(
                    "select count(*) from public.gasto where obra_id = %s", (ids["work"],)
                ).fetchone()[0]
                paid_advance = pay(advance.json()["id"])
                assert paid_advance["state"] == "pagado" and paid_advance["paid_by"] == "Test admin"

                # Financial bridge: the payment is one more validated expense, same transaction.
                assert connection.execute(
                    "select count(*) from public.gasto where obra_id = %s", (ids["work"],)
                ).fetchone()[0] == expenses_before + 1
                bridged = connection.execute(
                    """select g.id, g.folio, g.estado::text, g.importe, g.subtotal, g.iva,
                              g.proveedor_id, g.partida_gasto_id, g.subpartida_gasto_id,
                              cag.nombre, g.concepto, g.fecha,
                              (now() at time zone 'America/Mexico_City')::date, g.area_id,
                              g.validado_por, (select count(*) from public.gasto_concepto c
                                               where c.gasto_id = g.id)
                       from public.gasto g
                       join public.catalogo_categoria_gasto cag on cag.id = g.categoria_gasto_id
                       where g.estimacion_subcontrato_id = %s""",
                    (advance.json()["id"],),
                ).fetchone()
                assert bridged[2:10] == (
                    "validado", D("1000"), D("1000"), D("0"), ids["supplier"], ids["item"],
                    ids["subitem"], "MANO DE OBRA",
                )
                assert bridged[10] == f"Pago de Estimación EST-01 - Subcontrato {contract['folio']}"
                assert bridged[11] == bridged[12]  # payment day in Mexico
                assert bridged[13] is None and bridged[14] == ids["admin"] and bridged[15] == 1
                assert paid_advance["expense_id"] == str(bridged[0])
                assert paid_advance["expense_folio"] == bridged[1]
                # Managed from Subcontratos only: the expense endpoints refuse it.
                expense_url = f"/api/v1/expenses/{bridged[0]}"
                assert api.post(f"{expense_url}/review", json={"action": "return_to_review",
                                                          "reason": "Prueba de bloqueo"},
                                headers=headers("admin")).status_code == 409
                assert api.post(f"{expense_url}/cancel", json={"reason": "Prueba de bloqueo"},
                                headers=headers("admin")).status_code == 409

                progress = post({"kind": "avance", "gross_amount": "4000",
                                 "advance_amortization": "500", "additions": "200",
                                 "deductions": "100", "adjustment_notes": "Muro extra; daño"})
                assert progress.status_code == 201, progress.text
                p = progress.json()
                assert (p["folio"], p["retention_amount"], p["net_amount"]) == (
                    "EST-02", "200.0000", "3400.0000"
                )
                # Drafts can be edited or deleted; paid ones are immutable.
                edited = api.put(f"{est}/{p['id']}", json={
                    "estimated_on": "2026-09-21", "kind": "avance", "gross_amount": "4000",
                    "advance_amortization": "500", "additions": "200", "deductions": "100",
                    "adjustment_notes": "Muro extra; daño en puerta",
                }, headers=headers("admin"))
                assert edited.status_code == 200 and edited.json()["net_amount"] == "3400.0000"
                pay(p["id"])
                assert api.put(f"{est}/{p['id']}", json={
                    "estimated_on": "2026-09-21", "kind": "avance", "gross_amount": "1",
                }, headers=headers("admin")).status_code == 409
                assert api.delete(f"{est}/{p['id']}", headers=headers("admin")).status_code == 409
                receipt = f"{est}/{p['id']}/receipt.pdf"
                for role in ("admin", "operativo"):
                    pdf = api.get(receipt, headers=headers(role))
                    assert pdf.status_code == 200, pdf.text
                    assert pdf.headers["content-type"] == "application/pdf"
                    assert pdf.content.startswith(b"%PDF")
                    assert "recibo-SC-" in pdf.headers["content-disposition"]
                assert api.get(receipt, headers=headers("outsider")).status_code == 403
                assert api.get(f"{est}/{uuid4()}/receipt.pdf",
                               headers=headers("admin")).status_code == 404
                draft = post({"kind": "avance", "gross_amount": "10"})
                assert api.delete(
                    f"{est}/{draft.json()['id']}", headers=headers("admin")
                ).status_code == 204

                # The database enforces the net equation even against direct writes.
                with pytest.raises(psycopg.errors.CheckViolation):
                    connection.execute(
                        """insert into public.estimacion_subcontrato
                           (subcontrato_id, numero, folio, fecha, tipo, importe_bruto,
                            importe_neto, creado_por)
                           values (%s, 99, 'EST-99', '2026-09-22', 'avance', 100, 101, %s)""",
                        (sid, ids["admin"]),
                    )

                # Overview: the generated expenses are the validated spend (no double count);
                # the contract adds to "comprometido" only what is still to be paid.
                overview = api.get(
                    f"/api/v1/works/{ids['work']}/overview", headers=headers("operativo"),
                ).json()
                totals = overview["totals"]
                legacy_contract = D("5000.5") if installation == "upgrade" else D("0")
                assert D(str(totals["subcontract_paid"])) == D("4400")
                assert D(str(totals["validated"])) == D("125") + D("4400")
                assert D(str(totals["subcontract_committed"])) == D("10100") + legacy_contract
                # 125 legacy + 4400 paid expenses + (10100 − 4400) still owed + legacy contract.
                assert D(str(totals["committed"])) == D("125") + D("10100") + legacy_contract
                labor = {row["name"]: row for row in overview["by_category"]}["MANO DE OBRA"]
                assert D(str(labor["validated"])) >= D("4400")
                provider = {row["name"]: row for row in overview["by_provider"]}["Test"]
                assert D(str(provider["validated"])) == D("4400")
                assert provider["expense_count"] == 2

                # The finiquito must amortize the whole pending advance (500) and settles.
                over = post({"kind": "avance", "gross_amount": "6000.01"})
                assert over.status_code == 422 and "excede" in over.text
                short = post({"kind": "finiquito", "gross_amount": "6000",
                              "advance_amortization": "400"})
                assert short.status_code == 422
                final = post({"kind": "finiquito", "gross_amount": "6000",
                              "advance_amortization": "500"})
                assert final.status_code == 201, final.text
                assert (final.json()["retention_amount"], final.json()["net_amount"]) == (
                    "300.0000", "5200.0000"
                )
                # A closed payment week blocks the payment and creates no expense.
                close_id = connection.execute(
                    """insert into public.cierre_semanal
                         (obra_id, anio_iso, semana_iso, cerrado_por)
                       select %s, extract(isoyear from d)::int, extract(week from d)::int, %s
                       from (select (now() at time zone 'America/Mexico_City')::date as d) t
                       returning id""",
                    (ids["work"], ids["admin"]),
                ).fetchone()[0]
                blocked = api.patch(f"{est}/{final.json()['id']}/status", json={"state": "pagado"},
                                    headers=headers("admin"))
                assert blocked.status_code == 409 and "cerrada" in blocked.text
                assert connection.execute(
                    "select count(*) from public.gasto where estimacion_subcontrato_id = %s",
                    (final.json()["id"],),
                ).fetchone()[0] == 0
                connection.execute("delete from public.cierre_semanal where id = %s", (close_id,))
                pay(final.json()["id"])
                settled = api.get(f"/api/v1/subcontracts/{sid}", headers=headers("operativo"))
                detail = settled.json()
                assert detail["state"] == "finiquitado"
                assert (detail["paid_net"], detail["retained"],
                        detail["advance_pending_amortization"]) == (
                    "9600.0000", "500.0000", "0.0000"
                )
                assert [e["folio"] for e in detail["estimations"]] == ["EST-01", "EST-02", "EST-04"]
                assert post({"kind": "avance", "gross_amount": "1"}).status_code == 409
                assert api.patch(f"/api/v1/subcontracts/{sid}", json={"description": "Nuevo"},
                                 headers=headers("admin")).status_code == 409
                overview = api.get(
                    f"/api/v1/works/{ids['work']}/overview", headers=headers("admin"),
                ).json()
                # Settled: committed is the work actually valued (4000+200−100+6000).
                assert D(str(overview["totals"]["subcontract_committed"])) == (
                    D("10100") + legacy_contract
                )
                assert D(str(overview["totals"]["subcontract_paid"])) == D("9600")

                # Cancelling requires no drafts; afterwards only what was paid is committed.
                other = api.post(base, json={**body, "contracted_amount": "3000"},
                                 headers=headers("admin")).json()
                draft = api.post(f"/api/v1/subcontracts/{other['id']}/estimations", json={
                    "estimated_on": "2026-09-25", "kind": "avance", "gross_amount": "100",
                }, headers=headers("admin")).json()
                cancel = {"state": "cancelado"}
                assert api.patch(f"/api/v1/subcontracts/{other['id']}", json=cancel,
                                 headers=headers("admin")).status_code == 409
                api.delete(f"/api/v1/subcontracts/{other['id']}/estimations/{draft['id']}",
                           headers=headers("admin"))
                cancelled = api.patch(f"/api/v1/subcontracts/{other['id']}", json=cancel,
                                      headers=headers("admin"))
                assert cancelled.status_code == 200 and cancelled.json()["state"] == "cancelado"
                overview = api.get(
                    f"/api/v1/works/{ids['work']}/overview", headers=headers("admin"),
                ).json()
                assert D(str(overview["totals"]["subcontracted"])) == D("10000") + legacy_contract

                actions = [row[0] for row in connection.execute(
                    """select accion from public.audit_log_negocio
                       where entidad in ('subcontrato', 'estimacion_subcontrato')
                       order by creado_en, id"""
                ).fetchall()]
                assert {"crear", "editar", "pagar", "eliminar", "finiquitar", "cancelar"} <= set(
                    actions
                )
                # Weekly payroll (default: current week): worker, activity, folio, net + total.
                payroll_url = f"/api/v1/works/{ids['work']}/subcontracts/payroll-export"
                sheet_response = api.get(payroll_url, headers=headers("operativo"))
                assert sheet_response.status_code == 200, sheet_response.text
                assert sheet_response.headers["content-type"].startswith(
                    "application/vnd.openxmlformats"
                )
                sheet = openpyxl.load_workbook(io.BytesIO(sheet_response.content)).active
                values = [[cell.value for cell in row] for row in sheet.iter_rows()]
                assert values[2] == [
                    "Nombre del Trabajador", "Actividad Realizada", "Folio Pago", "Balance (Neto)"
                ]
                body_rows = values[3:-1]
                assert [row[2] for row in body_rows] == ["EST-01", "EST-02", "EST-04"]
                assert {row[0] for row in body_rows} == {"Test"}
                assert body_rows[0][1] == "Colocación de block en muros perimetrales"
                assert [D(str(row[3])) for row in body_rows] == [D("1000"), D("3400"), D("5200")]
                assert values[-1][2] == "TOTAL DE LA SEMANA"
                assert D(str(values[-1][3])) == D("9600")
                empty = api.get(payroll_url, params={"from": "2020-01-06", "to": "2020-01-12"},
                                headers=headers("admin"))
                empty_values = [[c.value for c in r] for r in openpyxl.load_workbook(
                    io.BytesIO(empty.content)).active.iter_rows()]
                assert len(empty_values) == 4 and D(str(empty_values[-1][3])) == D("0")
                assert api.get(payroll_url, params={"from": "2026-09-30", "to": "2026-09-01"},
                               headers=headers("admin")).status_code == 422
                assert api.get(payroll_url, headers=headers("outsider")).status_code == 403

                dashboard = api.get("/api/v1/dashboard", headers=headers("admin")).json()
                expected_paid = D("9600") + (D("1000") if installation == "upgrade" else D("0"))
                assert D(str(dashboard["totals"]["subcontract_paid"])) == expected_paid
        finally:
            app.dependency_overrides.clear()
            app.dependency_overrides.update(previous)


def test_bridge_migration_backfills_paid_estimations(isolated_services):  # noqa: F811
    """Estimations paid before the bridge get their validated expense, exactly once."""
    services = isolated_services
    with psycopg.connect(services["dsn"], autocommit=True) as admin:
        admin.execute("create database bridge_backfill")
    dsn = services["dsn"].removesuffix("postgres") + "bridge_backfill"
    with psycopg.connect(dsn, autocommit=True) as connection:
        connection.execute(BOOTSTRAP)
        for path in sorted(MIGRATIONS.glob("*.sql")):
            if path == BRIDGE:
                break
            connection.execute(path.read_text())
        ids = seed(connection)
        labor = connection.execute(
            "select id from public.catalogo_categoria_gasto where nombre = 'MANO DE OBRA'"
        ).fetchone()[0]
        contract = connection.execute(
            """insert into public.subcontrato (obra_id, proveedor_id, partida_gasto_id,
                 subpartida_gasto_id, categoria_gasto_id, descripcion, importe_contratado)
               values (%s, %s, %s, %s, %s, 'Aplanados', 10000) returning id, folio""",
            (ids["work"], ids["supplier"], ids["item"], ids["subitem"], labor),
        ).fetchone()
        for number, (state, net) in enumerate((("pagado", 2500), ("pagado", 0),
                                               ("borrador", 700)), 1):
            connection.execute(
                """insert into public.estimacion_subcontrato (subcontrato_id, numero, folio,
                     fecha, tipo, importe_bruto, importe_neto, estado, pagado_por, pagado_en,
                     creado_por, amortizacion_anticipo)
                   values (%s, %s, %s, '2026-09-20', 'avance', %s, %s, %s, %s,
                           case when %s = 'pagado'
                                then '2026-09-21 18:00+00'::timestamptz end, %s, %s)""",
                (contract[0], number, f"EST-{number:02d}", max(net, 100), net, state,
                 ids["admin"] if state == "pagado" else None, state, ids["admin"],
                 max(net, 100) - net),
            )
        for _ in range(2):  # idempotent
            connection.execute(BRIDGE.read_text())
        rows = connection.execute(
            """select e.folio, g.estado::text, g.importe, g.fecha, g.concepto
               from public.gasto g join public.estimacion_subcontrato e
                 on e.id = g.estimacion_subcontrato_id order by e.folio"""
        ).fetchall()
        # Only the paid estimation with money moved; drafts and zero nets create nothing.
        assert rows == [(
            "EST-01", "validado", D("2500"), date(2026, 9, 21),
            f"Pago de Estimación EST-01 - Subcontrato {contract[1]}",
        )]
