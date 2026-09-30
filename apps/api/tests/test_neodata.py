from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.neodata import (
    NEODATA_PARSER_VERSION,
    NeodataError,
    UnsafeWorkbookError,
    parse_neodata_workbook,
)


def workbook_bytes(*, external_link: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto"
    sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "P. Unitario", "Importe"])
    sheet.append([None, "PRESUPUESTO OFICINA", None, None, None, None])
    sheet.append(["PRELIMINARES", "PRELIMINARES", None, None, None, None])
    sheet.append(["PRE-01", "Trazo y nivelación", "M2", 2, 10, "=D4*E4"])
    sheet.append([None, "incluye herramienta", None, 0, 0, 0])
    sheet.append(["PRELIMINARES", "TOTAL PRELIMINARES", None, None, 0, 20])
    sheet.append(["TOTAL DEL PRESUPUESTO MOSTRADO SIN IVA:", None, None, None, None, 20])
    if external_link:
        sheet["H1"] = "='[external.xlsx]Sheet1'!A1"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_synthetic_preview_preserves_continuation_and_totals() -> None:
    preview = parse_neodata_workbook(workbook_bytes(), "presupuesto.xlsx")

    assert preview.areas == {"PRELIMINARES": 1}
    assert preview.items[0].area_path == ("Oficina", "PRELIMINARES")
    assert len(preview.items) == 1
    assert preview.items[0].description == "Trazo y nivelación incluye herramienta"
    assert str(preview.calculated_total_without_vat) == "20"
    assert preview.section_totals[0].difference == 0
    assert preview.unclassified == []


def test_generic_budget_promotes_repeated_headers_to_areas_and_skips_rollups() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "P. Unitario", "Importe"])
    sheet.append(["PRESUPUESTO", "PRESUPUESTO", None, None, None, None])
    sheet.append(["PRELIMINARES", "PRELIMINARES", None, None, None, None])
    sheet.append(["PRE-01", "Trazo", "M2", 3, 1.005, 3.02])
    sheet.append(["PRELIMINARES", "TOTAL PRELIMINARES", None, None, None, 3.02])
    sheet.append(["PRELIMINARES", "TOTAL PRELIMINARES", None, None, None, 3.02])
    sheet.append(["TOTAL DEL PRESUPUESTO MOSTRADO SIN IVA:", None, None, None, None, 3.02])
    output = BytesIO()
    workbook.save(output)

    preview = parse_neodata_workbook(output.getvalue(), "presupuesto.xlsx")

    assert preview.areas == {"PRELIMINARES": 1}
    assert preview.items[0].area_path == ("PRELIMINARES",)
    assert preview.calculated_total_without_vat == Decimal("3.02")
    assert len(preview.section_totals) == 1
    assert preview.rollup_total_count == 1
    assert preview.warnings == []


def test_named_budget_builds_hierarchy_and_closes_implicit_installations() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Código", "Concepto", "Unidad", "Cantidad", "P. Unitario", "Importe"])
    sheet.append([None, "PRESUPUESTO OFICINA", None, None, None, None])
    sheet.append(["INSTALACIONES", "INSTALACIONES", None, None, None, None])
    sheet.append(["ELECTRICAS", "ELECTRICAS", None, None, None, None])
    sheet.append(["ELE-01", "Contacto", "PZA", 2, 10, 20])
    sheet.append(["ELECTRICAS", "TOTAL ELECTRICAS", None, None, 0, 20])
    # NEODATA omite TOTAL INSTALACIONES antes de abrir el capítulo siguiente.
    sheet.append(["ACABADOS", "ACABADOS", None, None, None, None])
    sheet.append(["ACA-01", "Pintura", "M2", 3, 10, 30])
    sheet.append(["ACABADOS", "TOTAL ACABADOS", None, None, 0, 30])
    sheet.append(["TOTAL DEL PRESUPUESTO MOSTRADO SIN IVA:", None, None, None, None, 50])
    output = BytesIO()
    workbook.save(output)

    preview = parse_neodata_workbook(output.getvalue(), "presupuesto.xlsx")

    assert [item.area_path for item in preview.items] == [
        ("Oficina", "INSTALACIONES", "ELECTRICAS"),
        ("Oficina", "ACABADOS"),
    ]
    assert preview.to_dict()["parser_version"] == NEODATA_PARSER_VERSION
    assert preview.to_dict()["area_count"] == 2


@pytest.mark.parametrize(
    ("content", "filename", "error"),
    [
        (b"not-a-zip", "presupuesto.xlsx", NeodataError),
        (workbook_bytes(), "presupuesto.xlsm", UnsafeWorkbookError),
        (workbook_bytes(), "presupuesto.csv", NeodataError),
    ],
    ids=["corrupt", "macro", "wrong-extension"],
)
def test_rejects_unsafe_or_invalid_files(
    content: bytes, filename: str, error: type[Exception]
) -> None:
    with pytest.raises(error):
        parse_neodata_workbook(content, filename)


def test_rejects_external_formula_links() -> None:
    with pytest.raises(UnsafeWorkbookError, match="enlaces externos"):
        parse_neodata_workbook(workbook_bytes(external_link=True), "presupuesto.xlsx")


def test_enforces_file_size_limit() -> None:
    with pytest.raises(NeodataError, match="debe medir"):
        parse_neodata_workbook(workbook_bytes(), "presupuesto.xlsx", max_bytes=100)


def test_real_toluca_workbook_regression() -> None:
    fixture = Path(
        "/var/www/Apparquitectos/20260816 INFRA TOLUCA OFICINA-COMEDOR Y BAÑOS "
        "COMPETO JUNTOS (2).xlsx"
    )
    if not fixture.exists():
        pytest.skip("El Excel real se mantiene intencionalmente fuera de Git")

    preview = parse_neodata_workbook(fixture.read_bytes(), fixture.name)

    assert preview.sheets == ["b)Estandar (E)"]
    assert preview.row_count == 1154
    assert len(preview.area_tree) == 46
    assert len(preview.selectable_area_paths) == 40
    assert sum(item.area_path[0] == "Oficina" for item in preview.items) == 81
    assert sum(item.area_path[0] == "Comedor" for item in preview.items) == 65
    assert sum(item.area_path[0] == "Baños" for item in preview.items) == 86
    assert any(
        item.area_path == ("Oficina", "PRELIMINARES") for item in preview.items
    )
    assert any(
        item.area_path == ("Oficina", "INSTALACIONES", "ELECTRICAS")
        for item in preview.items
    )
    assert any(
        item.area_path == ("Comedor", "INSTALACIONES", "HIDRAULICAS")
        for item in preview.items
    )
    assert any(
        item.area_path == ("Baños", "DESMANTELACIONES") for item in preview.items
    )
    assert len(preview.items) == 232
    assert len(preview.section_totals) == 40
    assert preview.unclassified == []
    assert all(abs(total.difference) <= 0.02 for total in preview.section_totals)
    assert abs(
        preview.calculated_total_without_vat - Decimal("2624832.8848544")
    ) <= Decimal("0.001")


def test_real_casa_pse_workbook_regression() -> None:
    fixture = Path(
        "/var/www/Apparquitectos/20210621 Presupuesto Casa PSE COMPLETO (3).XLSX"
    )
    if not fixture.exists():
        pytest.skip("El Excel real se mantiene intencionalmente fuera de Git")

    preview = parse_neodata_workbook(fixture.read_bytes(), fixture.name)

    assert preview.sheets == ["b)Estandar (E)"]
    assert preview.row_count == 2918
    assert len(preview.areas) == 70
    assert sum(preview.areas.values()) == 520
    assert len(preview.area_tree) == 254
    assert any(
        item.area_path == ("CIMENTACION", "SOTANO") for item in preview.items
    )
    assert any(
        item.area_path == ("EXTRUCTURA", "ACERO", "PB", "COLUMNAS", "C-1")
        for item in preview.items
    )
    assert len(preview.items) == 520
    assert preview.consolidated_item_count == 520
    assert len(preview.section_totals) == 191
    assert preview.rollup_total_count == 68
    assert preview.unclassified == []
    assert all(abs(total.difference) <= Decimal("0.02") for total in preview.section_totals)
    assert preview.calculated_total_without_vat == Decimal("3882647.02")
