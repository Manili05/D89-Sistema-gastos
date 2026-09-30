import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

CATEGORY_NAMES = {"HIDRAULICAS", "SANITARIAS", "ELECTRICAS", "ESPECIALES"}
MONEY_TOLERANCE = Decimal("0.02")
NEODATA_PARSER_VERSION = 2


class NeodataError(ValueError):
    """Base error safe to expose to an administrator."""


class UnsafeWorkbookError(NeodataError):
    """Workbook contains a feature the importer deliberately refuses."""


def normalized_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text).strip().upper()


def decimal_value(value: object, default: Decimal = Decimal("0")) -> Decimal:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        raise NeodataError("Se encontró un booleano donde se esperaba un número")
    try:
        return Decimal(str(value))
    except Exception as exc:
        raise NeodataError(f"Valor numérico inválido: {value!r}") from exc


@dataclass
class BudgetItem:
    sheet: str
    row: int
    area: str
    work_class: str
    category: str | None
    code: str
    description: str
    unit: str
    quantity: Decimal
    unit_price: Decimal
    amount: Decimal
    area_path: tuple[str, ...] = ()

    @property
    def identity(self) -> tuple[str, str, str, str, str]:
        return (
            normalized_text(self.code),
            normalized_text(self.description),
            normalized_text(self.unit),
            normalized_text(self.work_class),
            normalized_text(self.category),
        )


@dataclass(frozen=True)
class SectionTotal:
    sheet: str
    row: int
    area: str
    work_class: str
    category: str | None
    label: str
    declared: Decimal
    calculated: Decimal
    area_path: tuple[str, ...] = ()

    @property
    def difference(self) -> Decimal:
        return self.calculated - self.declared


@dataclass(frozen=True)
class UnclassifiedRow:
    sheet: str
    row: int
    values: tuple[str, ...]
    reason: str


@dataclass
class ImportPreview:
    filename: str
    sheets: list[str] = field(default_factory=list)
    row_count: int = 0
    items: list[BudgetItem] = field(default_factory=list)
    section_totals: list[SectionTotal] = field(default_factory=list)
    rollup_total_count: int = 0
    unclassified: list[UnclassifiedRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    declared_total_without_vat: Decimal | None = None

    @property
    def calculated_total_without_vat(self) -> Decimal:
        return sum((item.amount for item in self.items), Decimal("0"))

    @property
    def areas(self) -> dict[str, int]:
        return dict(Counter(item.area for item in self.items))

    @property
    def selectable_area_paths(self) -> set[tuple[str, ...]]:
        return {item.area_path or (item.area,) for item in self.items}

    @property
    def area_tree(self) -> list[dict[str, object]]:
        paths: set[tuple[str, ...]] = set()
        selectable = self.selectable_area_paths
        for item_path in selectable:
            for depth in range(1, len(item_path) + 1):
                paths.add(item_path[:depth])
        return [
            {
                "name": path[-1],
                "path": list(path),
                "parent_path": list(path[:-1]) or None,
                "level": len(path) - 1,
                "selectable": path in selectable,
            }
            for path in sorted(paths, key=lambda value: (len(value), value))
        ]

    @property
    def consolidated_item_count(self) -> int:
        unique_items = {
            (
                tuple(normalized_text(part) for part in (item.area_path or (item.area,))),
                item.identity,
            )
            for item in self.items
        }
        return len(unique_items)

    def to_dict(self, include_items: bool = True) -> dict[str, object]:
        result: dict[str, object] = {
            "filename": self.filename,
            "parser_version": NEODATA_PARSER_VERSION,
            "sheets": self.sheets,
            "row_count": self.row_count,
            "area_count": len(self.selectable_area_paths),
            "area_label_count": len(self.areas),
            "areas": self.areas,
            "area_tree": self.area_tree,
            "item_count": len(self.items),
            "consolidated_item_count": self.consolidated_item_count,
            "section_total_count": len(self.section_totals),
            "rollup_total_count": self.rollup_total_count,
            "declared_total_without_vat": str(self.declared_total_without_vat or "0"),
            "calculated_total_without_vat": str(self.calculated_total_without_vat),
            "unclassified": [row.__dict__ for row in self.unclassified],
            "warnings": self.warnings,
            "section_totals": [
                {
                    "sheet": total.sheet,
                    "row": total.row,
                    "area": total.area,
                    "area_path": list(total.area_path or (total.area,)),
                    "work_class": total.work_class,
                    "category": total.category,
                    "label": total.label,
                    "declared": str(total.declared),
                    "calculated": str(total.calculated),
                    "difference": str(total.difference),
                }
                for total in self.section_totals
            ],
            "section_mismatches": [
                {
                    "sheet": total.sheet,
                    "row": total.row,
                    "label": total.label,
                    "difference": str(total.difference),
                }
                for total in self.section_totals
                if abs(total.difference) > MONEY_TOLERANCE
            ],
        }
        if include_items:
            result["items"] = [
                {
                    **item.__dict__,
                    "area_path": list(item.area_path or (item.area,)),
                    "quantity": str(item.quantity),
                    "unit_price": str(item.unit_price),
                    "amount": str(item.amount),
                    "identity": item.identity,
                }
                for item in self.items
            ]
        return result


def _first_area_from_title(title: str) -> str:
    raw = re.sub(r"^PRESUPUESTO\s+", "", title, flags=re.IGNORECASE)
    first = re.split(r",|\s+Y\s+", raw, maxsplit=1, flags=re.IGNORECASE)[0]
    return first.strip().title()


def _has_external_references(workbook: object) -> bool:
    if getattr(workbook, "_external_links", []):
        return True
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.hyperlink and str(cell.hyperlink.target).lower().startswith(
                    ("http:", "https:")
                ):
                    return True
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("=")
                    and re.search(r"\[[^\]]+\]", cell.value)
                ):
                    return True
    return False


def parse_neodata_workbook(
    content: bytes,
    filename: str,
    *,
    max_bytes: int = 10 * 1024 * 1024,
) -> ImportPreview:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsm", ".xltm"}:
        raise UnsafeWorkbookError("Los libros con macros no están permitidos")
    if suffix != ".xlsx":
        raise NeodataError("Solo se aceptan archivos .xlsx")
    if not content or len(content) > max_bytes:
        raise NeodataError(f"El archivo debe medir entre 1 byte y {max_bytes} bytes")
    if not content.startswith(b"PK"):
        raise NeodataError("El archivo no es un contenedor XLSX válido")

    try:
        formula_book = load_workbook(BytesIO(content), data_only=False, keep_links=True)
        if _has_external_references(formula_book):
            raise UnsafeWorkbookError("El libro contiene enlaces externos")
        formula_book.close()
        value_book = load_workbook(
            BytesIO(content), data_only=True, read_only=True, keep_links=False
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise NeodataError("No fue posible abrir el archivo XLSX") from exc

    preview = ImportPreview(filename=filename)
    for worksheet in value_book.worksheets:
        preview.sheets.append(worksheet.title)
        preview.row_count += worksheet.max_row
        rows = list(worksheet.iter_rows(values_only=True))
        current_area: str | None = None
        current_class: str | None = None
        current_category: str | None = None
        current_section_items: list[BudgetItem] = []
        last_item: BudgetItem | None = None
        budget_started = False
        generic_budget = False
        area_path: list[str] = []

        for index, raw_row in enumerate(rows, 1):
            values = list(raw_row) + [None] * 7
            a, b, unit, quantity, unit_price, declared_amount, _percent = values[:7]
            a_text = str(a).strip() if a is not None else ""
            b_text = str(b).strip() if b is not None else ""
            unit_text = str(unit).strip() if unit is not None else ""
            a_norm = normalized_text(a_text)
            b_norm = normalized_text(b_text)

            if a_norm.startswith("TOTAL DEL PRESUPUESTO MOSTRADO SIN IVA"):
                preview.declared_total_without_vat = decimal_value(declared_amount)
                budget_started = False
                last_item = None
                continue

            if b_norm == "PRESUPUESTO" or b_norm.startswith("PRESUPUESTO "):
                candidate = (
                    "General" if b_norm == "PRESUPUESTO" else _first_area_from_title(b_text)
                )
                if not budget_started:
                    budget_started = True
                    generic_budget = b_norm == "PRESUPUESTO"
                    current_area = candidate
                    area_path = [] if generic_budget else [candidate]
                else:
                    generic_budget = False
                    current_area = re.sub(
                        r"^PRESUPUESTO\s+", "", b_text, flags=re.IGNORECASE
                    ).strip().title()
                    area_path = [current_area]
                    current_class = current_category = None
                current_section_items.clear()
                last_item = None
                continue

            next_values = list(rows[index]) + [None] * 3 if index < len(rows) else [None] * 3
            next_a = normalized_text(next_values[0])
            next_b = normalized_text(next_values[1])
            is_area_header = (
                budget_started
                and a_norm
                and b_norm
                and a_norm != b_norm
                and not unit_text
                and quantity in (None, 0)
                and declared_amount in (None, 0)
                and next_a
                and next_a == next_b
            )
            if is_area_header:
                current_area = a_text.title()
                generic_budget = False
                area_path = [current_area]
                current_class = current_category = None
                current_section_items.clear()
                last_item = None
                continue

            if not budget_started or not current_area:
                continue

            if a_norm and b_norm.startswith("TOTAL "):
                total_path = (
                    tuple(area_path) if area_path else ((current_area,) if current_area else ())
                )
                if not current_section_items:
                    preview.rollup_total_count += 1
                    last_item = None
                    target = normalized_text(re.sub(r"^TOTAL\s+", "", b_text, flags=re.I))
                    match = next(
                        (position for position in range(len(area_path) - 1, -1, -1)
                         if normalized_text(area_path[position]) == target),
                        None,
                    )
                    if match is not None:
                        area_path = area_path[:match]
                        current_area = area_path[-1] if area_path else "General"
                    continue
                declared = decimal_value(declared_amount)
                calculated = sum((item.amount for item in current_section_items), Decimal("0"))
                total = SectionTotal(
                    sheet=worksheet.title,
                    row=index,
                    area=current_area,
                    work_class=current_class or "NO_CLASIFICADO",
                    category=current_category,
                    label=b_text,
                    declared=declared,
                    calculated=calculated,
                    area_path=total_path,
                )
                preview.section_totals.append(total)
                if abs(total.difference) > MONEY_TOLERANCE:
                    preview.warnings.append(
                        f"{worksheet.title}!{index}: {b_text} difiere por {total.difference}"
                    )
                current_section_items.clear()
                last_item = None
                target = normalized_text(re.sub(r"^TOTAL\s+", "", b_text, flags=re.I))
                match = next(
                    (position for position in range(len(area_path) - 1, -1, -1)
                     if normalized_text(area_path[position]) == target),
                    None,
                )
                if match is not None:
                    area_path = area_path[:match]
                    current_area = area_path[-1] if area_path else "General"
                continue

            if a_norm and b_norm and a_norm == b_norm and not unit_text:
                # Algunos reportes no imprimen "TOTAL INSTALACIONES" después de
                # cerrar sus subcapítulos. El siguiente capítulo no debe quedar
                # anidado accidentalmente bajo INSTALACIONES.
                if (
                    not generic_budget
                    and area_path
                    and normalized_text(area_path[-1]) == "INSTALACIONES"
                    and a_norm not in CATEGORY_NAMES
                ):
                    area_path.pop()

                area_path.append(a_text.strip())
                collapsed_path: list[str] = []
                for part in area_path:
                    if (
                        not collapsed_path
                        or normalized_text(collapsed_path[-1]) != normalized_text(part)
                    ):
                        collapsed_path.append(part)
                area_path = collapsed_path
                current_area = area_path[-1]
                class_position = 0 if generic_budget else min(1, len(area_path) - 1)
                current_class = normalized_text(area_path[class_position])
                current_category = (
                    a_text.title()
                    if current_class == "INSTALACIONES" and a_norm in CATEGORY_NAMES
                    else None
                )
                last_item = None
                continue

            is_item = (
                bool(a_text and b_text and unit_text)
                and isinstance(quantity, int | float)
                and not isinstance(quantity, bool)
                and isinstance(unit_price, int | float)
                and not isinstance(unit_price, bool)
            )
            if is_item:
                amount = (
                    decimal_value(declared_amount)
                    if declared_amount not in (None, "")
                    else decimal_value(quantity) * decimal_value(unit_price)
                )
                item = BudgetItem(
                    sheet=worksheet.title,
                    row=index,
                    area=current_area,
                    work_class=current_class or "NO_CLASIFICADO",
                    category=current_category,
                    code=a_text,
                    description=b_text,
                    unit=unit_text,
                    quantity=decimal_value(quantity),
                    unit_price=decimal_value(unit_price),
                    amount=amount,
                    area_path=tuple(
                        part
                        for position, part in enumerate(area_path or [current_area])
                        if position == 0
                        or normalized_text((area_path or [current_area])[position - 1])
                        != normalized_text(part)
                    ),
                )
                preview.items.append(item)
                current_section_items.append(item)
                last_item = item
                continue

            if not a_text and b_text and last_item is not None:
                last_item.description = f"{last_item.description} {b_text}".strip()
                continue

            if any(value not in (None, "", 0) for value in values[:7]):
                preview.unclassified.append(
                    UnclassifiedRow(
                        sheet=worksheet.title,
                        row=index,
                        values=tuple(str(value) for value in values[:7] if value not in (None, "")),
                        reason="fila activa sin patrón seguro",
                    )
                )
                last_item = None

    value_book.close()

    code_identities: dict[str, set[tuple[str, str, str, str, str]]] = defaultdict(set)
    for item in preview.items:
        code_identities[normalized_text(item.code)].add(item.identity)
    for code, identities in code_identities.items():
        if len(identities) > 1:
            preview.warnings.append(
                f"El código {code} aparece con {len(identities)} identidades compuestas; "
                "se conservarán por descripción, unidad, clase y categoría."
            )

    if preview.declared_total_without_vat is not None:
        difference = preview.calculated_total_without_vat - preview.declared_total_without_vat
        if abs(difference) > MONEY_TOLERANCE:
            preview.warnings.append(f"El total sin IVA difiere por {difference}")
    return preview
