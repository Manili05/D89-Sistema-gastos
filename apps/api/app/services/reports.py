from collections.abc import Iterable, Mapping
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _safe_excel_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def build_excel_report(
    title: str,
    columns: tuple[str, ...],
    rows: Iterable[Mapping[str, object]],
) -> bytes:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Reporte"
    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17233C")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.append(list(columns))
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C66A3D")
    for row in rows:
        sheet.append([_safe_excel_value(row.get(column, "")) for column in columns])
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, 1):
        values = [
            str(sheet.cell(row=row, column=index).value or "")
            for row in range(2, sheet.max_row + 1)
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max([len(column), *(len(value) for value in values)]) + 2,
            48,
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_pdf_report(
    title: str,
    columns: tuple[str, ...],
    rows: Iterable[Mapping[str, object]],
) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    data = [list(columns)]
    data.extend([[str(row.get(column, "")) for column in columns] for row in rows])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17233C")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8D3CB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F1EC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    document.build([Paragraph(title, styles["Title"]), Spacer(1, 12), table])
    return output.getvalue()


def _mxn(value: object) -> str:
    amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    sign = "-" if amount < 0 else ""
    return f"{sign}${abs(amount):,.2f}"


def build_estimation_receipt(data: Mapping[str, object]) -> bytes:
    """Payment receipt of one subcontract estimation, with a signature space.

    The breakdown follows the agreed formula
    Bruto + Aditivas − Deductivas − Amortización − Retención = Neto, from the stored
    (server-computed) amounts. A draft carries a visible "not valid" notice.
    """
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=44, bottomMargin=44,
        title=f"Recibo {data['subcontract_folio']} {data['estimation_folio']}",
    )
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=9, leading=12,
                           textColor=colors.HexColor("#4A5260"))
    story: list = [
        Paragraph("Recibo de pago de destajo", styles["Title"]),
        Paragraph(f"Obra: {escape(str(data['work_name']))}", small),
        Spacer(1, 10),
    ]
    if data["state"] != "pagado":
        story += [
            Paragraph(
                "<b>BORRADOR — sin validez como comprobante de pago</b>",
                ParagraphStyle("draft", parent=styles["BodyText"],
                               textColor=colors.HexColor("#B6493F"), fontSize=11),
            ),
            Spacer(1, 8),
        ]
    header = [
        ["Proveedor (destajista)", escape(str(data["supplier_name"] or "—"))],
        ["Subcontrato", f"{data['subcontract_folio']} · {escape(str(data['description'] or ''))}"],
        ["Partida / subpartida", escape(str(data["classification"] or "—"))],
        ["Estimación", f"{data['estimation_folio']} · {data['kind_label']}"],
        ["Fecha", str(data["estimated_on"])],
        ["Estado", "Pagada" + (f" el {data['paid_at']}" if data.get("paid_at") else "")
         if data["state"] == "pagado" else "Borrador"],
    ]
    info = Table([[Paragraph(f"<b>{k}</b>", small), Paragraph(v, small)] for k, v in header],
                 colWidths=[140, 330])
    info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2DDD4")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    lines = [
        ["Importe bruto", "", _mxn(data["gross_amount"])],
        ["+ Aditivas", "", _mxn(data["additions"])],
        ["− Deductivas", "", _mxn(data["deductions"])],
        ["− Amortización de anticipo", "", _mxn(data["advance_amortization"])],
        ["− Retención fondo de garantía", f"{data['retention_percent']} %",
         _mxn(data["retention_amount"])],
        ["= IMPORTE NETO A PAGAR", "", _mxn(data["net_amount"])],
    ]
    breakdown = Table(lines, colWidths=[260, 80, 130])
    breakdown.setStyle(TableStyle([
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#E2DDD4")),
        ("LINEABOVE", (0, -1), (-1, -1), 1.2, colors.HexColor("#17233C")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, -1), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [info, Spacer(1, 18), Paragraph("<b>Desglose</b>", styles["Heading4"]), breakdown]
    if data.get("adjustment_notes"):
        story += [Spacer(1, 10), Paragraph(
            f"<b>Notas de ajustes:</b> {escape(str(data['adjustment_notes']))}", small)]
    signatures = Table(
        [["", "", ""], ["Firma del Destajista", "", "Autorizó (D89)"],
         # Plain table cells are not markup: no escaping here (Paragraphs above are).
         [str(data["supplier_name"] or ""), "", ""]],
        colWidths=[210, 50, 210], rowHeights=[70, 16, 14],
    )
    signatures.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (0, 0), 0.8, colors.black),
        ("LINEBELOW", (2, 0), (2, 0), 0.8, colors.black),
        ("ALIGN", (0, 1), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
    ]))
    story += [Spacer(1, 36), signatures]
    document.build(story)
    return output.getvalue()
