from pathlib import Path
from io import BytesIO

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def load_env() -> dict[str, str]:
    return dict(
        line.split("=", 1)
        for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines()
        if line and not line.startswith("#") and "=" in line
    )


def main() -> None:
    env = load_env()
    with httpx.Client(base_url="http://127.0.0.1:3089", timeout=30) as client:
        login = client.post(
            "/auth/v1/token?grant_type=password",
            headers={"apikey": env["ANON_KEY"]},
            json={
                "email": env["STAGING_ADMIN_EMAIL"],
                "password": env["STAGING_ADMIN_PASSWORD"],
            },
        )
        login.raise_for_status()
        auth = {"Authorization": f"Bearer {login.json()['access_token']}"}
        dashboard = client.get("/api/v1/dashboard", headers=auth)
        dashboard.raise_for_status()
        payload = dashboard.json()
        works = payload["works"]
        totals = payload["totals"]
        assert works and str(totals["budget"]).startswith("2624832")
        assert totals["collected"] == "75000.0000"
        assert totals["subcontracted"] == "45000.0000"
        assert totals["subcontract_paid"] == "18450.0000"
        assert totals["cash_balance"] == "56550.0000"
        work_id = works[0]["id"]
        excel = client.get(f"/api/v1/reports/works/{work_id}.xlsx", headers=auth)
        excel.raise_for_status()
        assert excel.content.startswith(b"PK")
        workbook = load_workbook(filename=BytesIO(excel.content), read_only=True)
        assert workbook.active.max_row >= 3
        workbook.close()
        pdf = client.get(f"/api/v1/reports/works/{work_id}.pdf", headers=auth)
        pdf.raise_for_status()
        assert pdf.content.startswith(b"%PDF")
    print("Dashboard y exportaciones live aprobados: JSON, XLSX y PDF")


if __name__ == "__main__":
    main()
